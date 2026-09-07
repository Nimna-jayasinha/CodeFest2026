# script scans everything inside that folder, including subfolders.  saves it as documents.jsonl for later processing by the RAG/AI system.



from pathlib import Path
import json
import fitz
from docx import Document


# ---------------------------------------------
# PATHS
# ---------------------------------------------

CORPUS_DIR = Path(r"H:\Codefest2026\Ashen_Era_Archive")

OUTPUT_DIR = Path("data/processed")
OUTPUT_FILE = OUTPUT_DIR / "documents.jsonl"


# ---------------------------------------------
# TXT / MARKDOWN
# ---------------------------------------------

def read_text_file(path):
    try:
        text = path.read_text(
            encoding="utf-8",
            errors="ignore"
        )

        return [{
            "source": path.name,
            "path": str(path),
            "type": path.suffix.lower(),
            "page": None,
            "text": text
        }]

    except Exception as e:
        print(f"ERROR reading {path}: {e}")
        return []


# ---------------------------------------------
# PDF
# ---------------------------------------------

def read_pdf(path):
    pages = []

    try:
        document = fitz.open(path)

        for page_number, page in enumerate(document, start=1):

            text = page.get_text("text")

            pages.append({
                "source": path.name,
                "path": str(path),
                "type": ".pdf",
                "page": page_number,
                "text": text
            })

        document.close()

    except Exception as e:
        print(f"ERROR reading PDF {path}: {e}")

    return pages


# ---------------------------------------------
# DOCX
# ---------------------------------------------

def read_docx(path):
    try:
        document = Document(path)

        contents = []

        # Read normal paragraphs
        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                contents.append(text)

        # Read tables
        for table in document.tables:

            for row in table.rows:

                values = []

                for cell in row.cells:
                    values.append(cell.text.strip())

                contents.append(" | ".join(values))

        full_text = "\n".join(contents)

        return [{
            "source": path.name,
            "path": str(path),
            "type": ".docx",
            "page": None,
            "text": full_text
        }]

    except Exception as e:
        print(f"ERROR reading DOCX {path}: {e}")
        return []


# ---------------------------------------------
# DOCUMENT ROUTER
# ---------------------------------------------

def process_file(path):

    extension = path.suffix.lower()

    if extension == ".pdf":
        return read_pdf(path)

    elif extension == ".docx":
        return read_docx(path)

    elif extension in [".txt", ".md"]:
        return read_text_file(path)

    else:
        return []


# ---------------------------------------------
# MAIN
# ---------------------------------------------

def main():

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    records = []

    supported_files = 0
    skipped_images = 0

    print("Starting corpus ingestion...")
    print()

    for path in CORPUS_DIR.rglob("*"):

        if not path.is_file():
            continue

        extension = path.suffix.lower()

        if extension in [".pdf", ".docx", ".txt", ".md"]:

            supported_files += 1

            print(f"Reading: {path.name}")

            file_records = process_file(path)

            records.extend(file_records)

        elif extension == ".png":

            skipped_images += 1

    print()
    print("Saving extracted content...")

    with open(
        OUTPUT_FILE,
        "w",
        encoding="utf-8"
    ) as f:

        for record in records:

            json.dump(
                record,
                f,
                ensure_ascii=False
            )

            f.write("\n")

    print()
    print("================================")
    print("INGESTION COMPLETE")
    print("================================")

    print(f"Files processed: {supported_files}")
    print(f"Records created: {len(records)}")
    print(f"PNG files skipped: {skipped_images}")

    total_characters = sum(
        len(record["text"])
        for record in records
    )

    print(f"Characters extracted: {total_characters:,}")

    empty_records = sum(
        1
        for record in records
        if not record["text"].strip()
    )

    print(f"Empty records: {empty_records}")

    print()
    print(f"Saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()